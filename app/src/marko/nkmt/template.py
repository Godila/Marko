"""Шаблон выгрузки xlsx: шапка, строки-примеры и «Инструкция» — из parse.SPEC.

Единый источник колонок: рассинхрон шаблона с парсером невозможен по
построению (шапка листа = [s.title for s in SPEC]). Форматирование читабельное:
закреплённая шапка, жирные заголовки, ширины колонок по контенту, перенос
длинных описаний в «Инструкции».
"""
import io
import math

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from marko.nkmt.parse import SPEC, SPEC_SETS

# строки-примеры по ключам SPEC учат валидным значениям живых справочников ЧЗ:
# пол — литералы («ЖЕНСКИЙ», «УНИВЕРСАЛЬНЫЙ (УНИСЕКС)»), у шапок размер — из
# справочника 46–62 с системой ОБХВАТ ГОЛОВЫ (цвет — точные значения справочника)
EXAMPLES = [
    {"article": "AB-1001", "tnved": "6109100000", "name": "Футболка хлопковая",
     "product_type": "ФУТБОЛКА", "color": "БЕЛЫЙ", "composition": "100% хлопок",
     "size": "M", "size_system": "МЕЖДУНАРОДНЫЙ", "target_gender": "ЖЕНСКИЙ",
     "model": "AB-1001-M"},
    {"article": "GH-2001", "tnved": "6505009000", "name": "Шапка Мокко",
     "product_type": "ШАПКА", "color": "ОЛИВКОВЫЙ", "composition": "50% шерсть 50% акрил",
     "size": "54", "size_system": "ОБХВАТ ГОЛОВЫ", "target_gender": "УНИВЕРСАЛЬНЫЙ (УНИСЕКС)",
     "model": "GH-2001"},
]

# правила игры для листа «Инструкция»: метка + текст в колонке описания
NOTES = [
    ("Подстановка", "Порядок: значение из файла > правило РД (бренд × вид товара) > "
                    "дефолты из «Справочников». Правило и дефолт заполняют только "
                    "пустые ячейки строки."),
    ("Повторный импорт", "Артикул из прошлых батчей обновляется на месте; дубль "
                         "артикула внутри файла — ошибка строки."),
    ("Ошибки", "Видны в предпросмотре до импорта и в карточке после."),
    ("Техрегламент", "Системный: подставляется всегда из дефолтов консоли, "
                     "колонки в файле нет."),
]

INFO_WIDTHS = {"A": 22, "B": 14, "C": 16, "D": 84}  # Колонка/Обязательная/Подставляется/Описание
_D_COL_CHARS = INFO_WIDTHS["D"] - 4                  # символов в строке переноса


def build_template() -> bytes:
    """Лист «Выгрузка» (шапка + примеры) и лист «Инструкция» (колонка,
    обязательность, подстановка, описание + правила игры)."""
    return _build_book(SPEC, "Выгрузка", EXAMPLES, NOTES)


# примеры наборов: артикулы компонентов — ссылки на карточки из каталога,
# количество после «x» (умолчание 1); вторая строка — набор без привязки
EXAMPLES_SETS = [
    {"article": "SET-0001", "name": "", "tnved": "6505009000",
     "components": "GH-2001; SHARF-01x2", "composition": "подарочная упаковка"},
    {"article": "SET-0002", "name": "Набор из 2 предметов", "tnved": "6505009000",
     "count": 2},
]
NOTES_SETS = [
    ("Компоненты", "Артикул нашей карточки (приоритет; работает до генерации её "
                   "GTIN) или внешний GTIN 13–14 цифр, количество после ×/x, "
                   "умолчание 1. Разделитель — точка с запятой."),
    ("Набор без привязки", "Пустая колонка «Компоненты» + «Кол-во предметов» — "
                           "лёгпром допускает набор-«количество»: в КИН можно "
                           "складывать любые товары в этом количестве."),
    ("Ещё не опубликованные", "Компонент-черновик — предупреждение в предпросмотре: "
                              "набор сохранится черновиком, подача фида подождёт "
                              "публикации компонента."),
    ("Анти-дубли", "Повтор строки с тем же артикулом обновляет набор (не дублирует); "
                   "точный дубль состава другим артикулом — ошибка."),
    ("После подачи", "Состав поданного набора не меняется (КИН должен "
                     "соответствовать карточке) — соберите аналог."),
]


def build_set_template() -> bytes:
    """Шаблон импорта наборов: листы «Наборы» и «Инструкция»."""
    return _build_book(SPEC_SETS, "Наборы", EXAMPLES_SETS, NOTES_SETS)


def _build_book(spec: list, sheet: str, examples: list[dict], notes: list) -> bytes:
    wb = openpyxl.Workbook()
    head_font = Font(bold=True)

    ws = wb.active
    ws.title = sheet
    cols = [s for s in spec if s.title]
    ws.append([s.title for s in cols])
    for ex in examples:
        ws.append([ex.get(s.key, "") for s in cols])
    for cell in ws[1]:
        cell.font = head_font
        cell.alignment = Alignment(vertical="center")
    for i, s in enumerate(cols, start=1):
        longest = max([len(s.title)]
                      + [len(str(ex.get(s.key, ""))) for ex in examples])
        ws.column_dimensions[get_column_letter(i)].width = min(34, max(12, longest * 0.9 + 4))
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Инструкция")
    info.append(["Колонка", "Обязательная", "Подставляется", "Описание"])
    for s in spec:
        info.append([s.title, "да" if s.required else "нет",
                     "да" if s.defaultable else "—", s.hint or "—"])
    for label, text in notes:
        info.append([label, "", "", text])
    for cell in info[1]:
        cell.font = head_font
    for col, width in INFO_WIDTHS.items():
        info.column_dimensions[col].width = width
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in info.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    for r in range(2, info.max_row + 1):
        text = str(info.cell(row=r, column=4).value or "")
        info.row_dimensions[r].height = max(15, 13.5 * math.ceil(
            max(len(text), 1) / _D_COL_CHARS))
    info.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
