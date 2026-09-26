"""Наборы (sets) НК: конструктор + импорт xlsx + гард подачи + анти-дубли.

Спека: crpt-specs/sets-design-2026-09-22.md. Атрибутный состав набора —
live 22.09 (2478/2504/23821/16271, без декларации/вида/цвета).
"""
import io

import openpyxl
import pytest

from marko.nkmt.models import Batch, BrandCache, Card, SetItem
from marko.nkmt.parse import SPEC_SETS
from tests.test_api_nkmt_dicts import AUTH, AUTH_RO, client  # noqa: F401

TN = "6505009000"


@pytest.fixture
def nk(db, monkeypatch):
    """Сеть НК подменена: категории — одна, бренд — из кэша, токен константа.
    product/generate_gtins/feed мокаются точечно по тестам."""
    from marko.nkmt.client import NkClient
    monkeypatch.setattr("marko.connector_mt.manager.get_token", lambda _db: "T")
    monkeypatch.setattr(NkClient, "categories",
                        lambda self, token, tnved: [{"cat_id": 30728,
                                                      "cat_name": "Головные уборы"}])
    db.add(BrandCache(name="ycpb", brand_id=2102811))
    db.commit()


def _card(db, article, gtin="", status="published", name=None, tnved=TN,
          is_set=False):
    b = Batch(source_filename="seed", status="new")
    db.add(b); db.flush()
    c = Card(batch_id=b.id, article=article, gtin=gtin, tnved=tnved,
             name=name or article, status=status, is_set=is_set)
    db.add(c); db.commit()
    return c


def _xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([s.title for s in SPEC_SETS])
    for r in rows:
        ws.append([r.get(s.key, "") for s in SPEC_SETS])
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def _comp(article, quantity=1):
    return {"ref": article, "quantity": quantity}


# --- парсер packed-колонки ---

def test_parse_components():
    from marko.nkmt.parse import parse_components
    assert parse_components("AB-1; 0463056232235×2; SHARF x3") == [
        {"ref": "AB-1", "quantity": 1},
        {"ref": "0463056232235", "quantity": 2},
        {"ref": "SHARF", "quantity": 3}]
    # известный артикул целиком не режется количеством
    known = lambda a: a == "HX2"  # noqa: E731
    assert parse_components("HX2", known_article=known) == [{"ref": "HX2", "quantity": 1}]
    assert parse_components("HX2") == [{"ref": "H", "quantity": 2}]
    assert parse_components("") == []
    assert parse_components("A×0") == [{"ref": "A", "quantity": 0}]  # qty<1 — ошибка строки в build


# --- превью конструктора ---

def test_sets_preview_bound_ok(db, client, nk):
    _card(db, "HAT", gtin="04630562322355", name="Шапка Мокко")
    _card(db, "SCARF", gtin="04630562322356", name="Шарф Серый")
    r = client.post("/v1/nkmt/sets/preview", headers=AUTH, json={
        "components": [_comp("HAT"), _comp("SCARF")]}).json()
    assert r["ok"] is True and r["error"] == ""
    assert r["mode"] == "bound" and r["count"] == 2
    assert r["name"] == "Набор: Шапка Мокко + Шарф Серый"  # авто-наименование
    assert r["brand"] == "YCPB"                            # дефолтный бренд
    assert r["tnved"] == TN                                # из первого компонента
    assert [(c["article"], c["kind"], c["status"]) for c in r["components"]] \
        == [("HAT", "ours", "published"), ("SCARF", "ours", "published")]
    assert r["warnings"]  # ТНВЭД-подстановка предупреждает


def test_sets_preview_draft_component_warns_not_blocks(db, client, nk):
    """«Компонент не опубликован» — предупреждение: черновик набора создать
    можно, подача фида подождёт публикации компонента."""
    _card(db, "DRAFT", gtin="04630562322357", status="ok", name="Шапка Чертёж")
    r = client.post("/v1/nkmt/sets/preview", headers=AUTH, json={
        "components": [_comp("DRAFT")]}).json()
    assert r["ok"] is True
    assert any("не опубликован" in w for w in r["warnings"])


def test_sets_preview_unknown_dup_self_and_nested(db, client, nk, monkeypatch):
    from marko.nkmt.client import NkClient
    from marko.nkmt.client import NkHttpError
    _card(db, "HAT", gtin="04630562322355", name="Шапка")
    _card(db, "SETX", gtin="04630562322358", name="Другой набор", is_set=True)
    # неизвестная ссылка — не артикул и не GTIN
    r = client.post("/v1/nkmt/sets/preview", headers=AUTH, json={
        "article": "S1", "components": [_comp("абв")], "tnved": TN,
        "name": "Н"}).json()
    assert r["ok"] is False and "не похоже" not in r["error"] and r["error"]
    # дубль ссылки
    r = client.post("/v1/nkmt/sets/preview", headers=AUTH, json={
        "article": "S1", "components": [_comp("HAT"), _comp("hat")],
        "tnved": TN, "name": "Н"}).json()
    assert r["ok"] is False and "дважды" in r["error"]
    # компонент сам набор
    r = client.post("/v1/nkmt/sets/preview", headers=AUTH, json={
        "article": "S1", "components": [_comp("SETX")], "tnved": TN,
        "name": "Н"}).json()
    assert r["ok"] is False and "сам набор" in r["error"]
    # внешний GTIN: найден в НК и не набор — ок; сам набор — блок; нет — warning
    calls = {"04630562322399": {"is_set": False, "good_name": "Шарф из НК",
                                "good_status": "published"},
             "04630562322398": {"is_set": True},
             "04630562322397": None}   # None → 404
    def fake_product(self, token, gtin):
        val = calls.get(gtin)
        if val is None:
            raise NkHttpError(404, "not found")
        return val
    monkeypatch.setattr(NkClient, "product", fake_product)
    ok = client.post("/v1/nkmt/sets/preview", headers=AUTH, json={
        "article": "S1", "components": [_comp("04630562322399")],
        "tnved": TN, "name": "Н"}).json()
    assert ok["ok"] is True and ok["components"][0]["kind"] == "external"
    assert ok["components"][0]["name"] == "Шарф из НК"
    bad = client.post("/v1/nkmt/sets/preview", headers=AUTH, json={
        "article": "S1", "components": [_comp("04630562322398")],
        "tnved": TN, "name": "Н"}).json()
    assert bad["ok"] is False and "сам набор" in bad["error"]
    gone = client.post("/v1/nkmt/sets/preview", headers=AUTH, json={
        "article": "S1", "components": [_comp("04630562322397")],
        "tnved": TN, "name": "Н"}).json()
    assert gone["ok"] is True and any("не найден" in w for w in gone["warnings"])


def test_sets_unbound_mode(db, client, nk):
    """Набор без привязки: только количество (лёгпром допускает)."""
    r = client.post("/v1/nkmt/sets/preview", headers=AUTH, json={
        "article": "S2", "count": 2, "tnved": TN, "name": "Набор из 2"}).json()
    assert r["ok"] is True and r["mode"] == "unbound" and r["count"] == 2
    assert r["components"] == []
    r = client.post("/v1/nkmt/sets/preview", headers=AUTH, json={
        "article": "S3", "tnved": TN, "name": "Н"}).json()
    assert r["ok"] is False and "компоненты или количество" in r["error"]


# --- CRUD конструктора + анти-дубли ---

def test_sets_create_list_update_delete(db, client, nk):
    _card(db, "HAT", gtin="04630562322355", name="Шапка")
    _card(db, "SCARF", gtin="04630562322356", name="Шарф")
    body = {"components": [_comp("HAT"), _comp("SCARF")]}
    r = client.post("/v1/nkmt/sets", headers=AUTH, json=body).json()
    assert r["id"] and r["batch_id"]
    rows = client.get("/v1/nkmt/sets", headers=AUTH_RO).json()
    assert len(rows) == 1
    s = rows[0]
    assert s["article"] == "SET-0001"          # автоартикул конструктора
    assert s["mode"] == "bound" and s["count"] == 2
    assert [(c["article"], c["quantity"], c["kind"]) for c in s["components"]] \
        == [("HAT", 1, "ours"), ("SCARF", 1, "ours")]
    # точный дубль состава — блок (даже другим артикулом)
    dup = client.post("/v1/nkmt/sets/preview", headers=AUTH, json=body).json()
    assert dup["ok"] is False and "дублирует набор SET-0001" in dup["error"]
    # конструктор не обновляет чужой набор молча — артикул занят
    clash = client.post("/v1/nkmt/sets", headers=AUTH, json={
        "article": "SET-0001", "components": [_comp("HAT")],
        "tnved": TN, "name": "Н"})
    assert clash.status_code == 400 and "уже используется набором" in clash.json()["detail"]
    # другая пропорция — не дубль
    other = client.post("/v1/nkmt/sets/preview", headers=AUTH, json={
        "components": [_comp("HAT", 2), _comp("SCARF")]}).json()
    assert other["ok"] is True
    # правка до подачи: полная замена состава
    upd = client.put(f"/v1/nkmt/sets/{r['id']}", headers=AUTH, json={
        "components": [_comp("HAT", 2)]}).json()
    assert upd["id"] == r["id"]
    assert client.get("/v1/nkmt/sets", headers=AUTH_RO).json()[0]["count"] == 2
    # RO-скопу — 403
    assert client.put(f"/v1/nkmt/sets/{r['id']}", headers=AUTH_RO,
                      json=body).status_code == 403
    assert client.post("/v1/nkmt/sets", headers=AUTH_RO,
                       json=body).status_code == 403
    # удаление черновика: карточка, компоненты и пустой батч
    bid = client.get("/v1/nkmt/sets", headers=AUTH_RO).json()[0]["batch_id"]
    assert client.delete(f"/v1/nkmt/sets/{r['id']}", headers=AUTH).json() == {"ok": True}
    assert client.get("/v1/nkmt/sets", headers=AUTH_RO).json() == []
    assert db.query(SetItem).count() == 0
    assert db.get(Batch, bid) is None
    assert client.delete("/v1/nkmt/sets/999", headers=AUTH).status_code == 404


def test_sets_update_after_feed_409(db, client, nk):
    _card(db, "HAT", gtin="04630562322355", name="Шапка")
    sid = client.post("/v1/nkmt/sets", headers=AUTH,
                      json={"components": [_comp("HAT")]}).json()["id"]
    db.query(Card).filter(Card.id == sid).update({"status": "fed"})
    db.commit()
    assert client.put(f"/v1/nkmt/sets/{sid}", headers=AUTH,
                      json={"components": [_comp("HAT")]}).status_code == 409
    assert client.delete(f"/v1/nkmt/sets/{sid}", headers=AUTH).status_code == 409


# --- гард подачи и entry фида ---

class FakeFeed:
    def __init__(self):
        self.fed = []
        self.calls = []

    def categories(self, token, tnved):
        return [{"cat_id": 30728, "cat_name": "Головные уборы"}]

    def generate_gtins(self, token, quantity):
        return {"drafts": [{"gtin": f"4630520699{i:04d}"} for i in range(quantity)],
                "monthly-limit": {}}

    def feed(self, token, entries):
        self.fed.append(entries)
        return {"feed_id": 42}

    def product(self, token, gtin):
        self.calls.append(gtin)
        return {"is_set": False, "good_name": "внешний", "good_status": "published"}


def test_sets_feed_guard_and_entry(db, client, nk, monkeypatch):
    draft = _card(db, "HAT", gtin="", status="ok", name="Шапка")  # без gtin, черновик
    ext = "04630562322399"
    fake = FakeFeed()
    monkeypatch.setattr("marko.nkmt.client.NkClient", lambda base: fake)
    sid = client.post("/v1/nkmt/sets", headers=AUTH, json={
        "components": [_comp("HAT"), _comp(ext)]}).json()["id"]
    # гард: компонент не опубликован и без GTIN — блок с перечнем
    blocked = client.post(f"/v1/nkmt/sets/{sid}/feed", headers=AUTH)
    assert blocked.status_code == 409 and "не опубликован" in blocked.json()["detail"]
    # публикуем компонент (GTIN появится) — подача проходит
    db.query(Card).filter(Card.id == draft.id).update(
        {"status": "published", "gtin": "04630562322355"})
    db.commit()
    out = client.post(f"/v1/nkmt/sets/{sid}/feed", headers=AUTH).json()
    assert out == {"feed_id": 42, "feed_ids": [42]}
    s = next(x for x in client.get("/v1/nkmt/sets", headers=AUTH_RO).json()
             if x["id"] == sid)
    assert s["status"] == "fed" and s["gtin"] == "46305206990000"
    e = fake.fed[0][0]
    assert e["is_set"] is True and e["moderation"] == 1 and e["categories"] == [30728]
    assert e["set_gtins"] == [{"gtin": "04630562322355", "quantity": 1},
                              {"gtin": ext, "quantity": 1}]
    ga = {a["attr_id"]: a["attr_value"] for a in e["good_attrs"]}
    assert ga == {2478: s["name"], 23821: 2}  # 2504 → entry.brand, 16271 пуст


def test_feed_entry_set_resolves_article_src(db):
    """entry набора резолвит GTIN компонентов по article_src на момент подачи
    (компонент мог получить GTIN после сборки набора)."""
    from marko.nkmt.service import _feed_entry
    b = Batch(source_filename="x")
    db.add(b); db.flush()
    comp = Card(batch_id=b.id, article="C1", gtin="04630562322355", tnved=TN,
                name="c", status="published")
    db.add(comp)
    s = Card(batch_id=b.id, article="SET-9", gtin="04630562322360", tnved=TN,
             name="набор", cat_id="30728", status="ok", is_set=True,
             attributes={"2478": "набор", "2504": "YCPB", "23821": 1, "16271": ""})
    db.add(s); db.flush()
    db.add(SetItem(card_id=s.id, gtin="", article_src="C1", quantity=1))  # gtin тогда не было
    db.commit()
    e = _feed_entry(s, db)
    assert e["set_gtins"] == [{"gtin": "04630562322355", "quantity": 1}]
    assert e["brand"] == "YCPB" and e["is_set"] is True
    # непривязанный набор — без set_gtins
    db.query(SetItem).delete()
    db.commit()
    assert "set_gtins" not in _feed_entry(s, db)


# --- импорт xlsx: идемпотентность и защита артикулов ---

def test_sets_import_preview_and_idempotency(db, client, nk, monkeypatch):
    _card(db, "HAT", gtin="04630562322355", name="Шапка")
    fake = FakeFeed()  # product для внешнего GTIN — без живой сети
    monkeypatch.setattr("marko.nkmt.client.NkClient", lambda base: fake)
    data = _xlsx([
        {"article": "SET-A", "components": "HAT; 04630562322399"},
        {"article": "SET-B", "count": 2, "tnved": TN, "name": "Набор из 2"},
    ])
    pv = client.post("/v1/nkmt/sets/import/preview", headers=AUTH,
                     files={"file": ("sets.xlsx", data)}).json()
    assert pv["stats"]["ok"] == 2 and pv["stats"]["error"] == 0
    assert pv["rows"][0]["mode"] == "bound" and pv["rows"][0]["count"] == 2
    assert pv["rows"][1]["mode"] == "unbound"
    imp = client.post("/v1/nkmt/sets/import", headers=AUTH,
                      files={"file": ("sets.xlsx", data)}).json()
    assert imp["stats"]["ok"] == 2
    assert len(client.get("/v1/nkmt/sets", headers=AUTH_RO).json()) == 2
    assert db.query(SetItem).count() == 2
    # повторный импорт того же файла — обновление на месте, не дубли
    again = client.post("/v1/nkmt/sets/import", headers=AUTH,
                        files={"file": ("sets.xlsx", data)}).json()
    assert again["stats"]["ok"] == 2 and again["batch_id"] != imp["batch_id"]
    rows = client.get("/v1/nkmt/sets", headers=AUTH_RO).json()
    assert len(rows) == 2 and db.query(SetItem).count() == 2
    # превью = импорт бит-в-бит: дубль-состав строкой «SET-A2» с тем же составом
    dup_data = _xlsx([{"article": "SET-A2", "components": "HAT; 04630562322399"}])
    pv2 = client.post("/v1/nkmt/sets/import/preview", headers=AUTH,
                      files={"file": ("sets.xlsx", dup_data)}).json()
    assert pv2["rows"][0]["ok"] is False and "дублирует" in pv2["rows"][0]["error"]


def test_sets_import_article_collision(db, client, nk):
    """Артикул набора не может занять артикул обычной карточки: строка —
    ошибка, чужая карточка не тронута."""
    hat = _card(db, "HAT", gtin="04630562322355", name="Шапка")
    data = _xlsx([{"article": "HAT", "components": "HAT", "tnved": TN,
                   "name": "злой набор"}])
    imp = client.post("/v1/nkmt/sets/import", headers=AUTH,
                      files={"file": ("sets.xlsx", data)}).json()
    assert imp["stats"]["error"] == 1
    db.refresh(hat)
    assert hat.is_set is False and hat.name == "Шапка" and hat.status == "published"
    assert db.query(Card).filter(Card.is_set.is_(True)).count() == 0


def test_sets_reimport_of_submitted_set_is_rejected(db, client, nk, monkeypatch):
    """P1-регресс: повторный импорт своей строки ПОСЛЕ подачи не сбрасывает
    поданный набор в черновик — строка ошибочна и не пишется вовсе."""
    _card(db, "HAT", gtin="04630562322355", name="Шапка")
    fake = FakeFeed()
    monkeypatch.setattr("marko.nkmt.client.NkClient", lambda base: fake)
    data = _xlsx([{"article": "SET-A", "components": "HAT"}])
    client.post("/v1/nkmt/sets/import", headers=AUTH,
                files={"file": ("sets.xlsx", data)})
    card = db.query(Card).filter_by(article="SET-A").one()
    batch_before, items_before = card.batch_id, db.query(SetItem).count()
    db.query(Card).filter(Card.id == card.id).update({"status": "fed"})
    db.commit()
    again = client.post("/v1/nkmt/sets/import", headers=AUTH,
                        files={"file": ("sets.xlsx", data)}).json()
    assert again["stats"]["error"] == 1
    db.refresh(card)
    assert card.status == "fed" and card.batch_id == batch_before
    assert db.query(SetItem).count() == items_before  # состав не заменён
    # конструктор с артикулом поданного набора — 400, не перезапись
    r = client.post("/v1/nkmt/sets", headers=AUTH,
                    json={"article": "SET-A", "components": [_comp("HAT")],
                          "tnved": TN, "name": "Н"})
    assert r.status_code == 400 and "уже подан" in r.json()["detail"]


# --- пикер и проверка в ЧЗ ---

def test_sets_cards_picker(db, client, nk):
    _card(db, "HAT", gtin="04630562322355", name="Шапка Серая")
    _card(db, "DRAFT", gtin="04630562322357", status="ok", name="Чертёж")
    _card(db, "NOGTIN", gtin="", name="Без GTIN")
    rows = client.get("/v1/nkmt/sets/cards", headers=AUTH_RO).json()
    assert [r["article"] for r in rows] == ["HAT"]  # только published с GTIN
    q = client.get("/v1/nkmt/sets/cards?q=сер", headers=AUTH_RO).json()
    assert [r["article"] for r in q] == ["HAT"]


def test_sets_check_in_nk(db, client, nk, monkeypatch):
    from marko.nkmt.client import NkClient
    _card(db, "HAT", gtin="04630562322355", name="Шапка")
    sid = client.post("/v1/nkmt/sets", headers=AUTH,
                      json={"components": [_comp("HAT")]}).json()["id"]
    assert client.post(f"/v1/nkmt/sets/{sid}/check", headers=AUTH).status_code == 409
    db.query(Card).filter(Card.id == sid).update({"gtin": "04630562322360"})
    db.commit()
    monkeypatch.setattr(NkClient, "product", lambda self, token, gtin: {
        "good_status": "published", "good_name": "Набор: Шапка",
        "is_set": True, "set_gtins": [{"gtin": "04630562322355", "quantity": 1}],
        "update_date": "2026-09-26"})
    out = client.post(f"/v1/nkmt/sets/{sid}/check", headers=AUTH).json()
    assert out == {"found": True, "good_status": "published",
                   "name": "Набор: Шапка", "is_set": True,
                   "set_gtins": [{"gtin": "04630562322355", "quantity": 1}],
                   "update_date": "2026-09-26"}


# --- кэш атрибутов набора и шаблон ---

def test_attrs_model_is_set_cache(db):
    from marko.nkmt import dicts
    from marko.platform.models import PlatformKV

    class FakeAttrs:
        def __init__(self):
            self.calls = []

        def attributes(self, token, tnved, attr_type=None, is_set=False):
            self.calls.append((tnved, attr_type, is_set))
            return [{"attr_id": 23821 if is_set else 12}]

    fake = FakeAttrs()
    out = dicts.attrs_model(db, fake, "T", TN, is_set=True)
    assert out == {"m": [{"attr_id": 23821}], "r": [{"attr_id": 23821}]}
    assert all(c[2] is True for c in fake.calls) and len(fake.calls) == 2
    kv = db.get(PlatformKV, f"nk_attrs_set:{TN}")  # отдельный ключ кэша
    assert kv is not None
    dicts.attrs_model(db, fake, "T", TN, is_set=True)  # из кэша, без сети
    assert len(fake.calls) == 2


def test_sets_template_endpoint(db, client):
    r = client.get("/v1/nkmt/sets/import/template", headers=AUTH_RO)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).worksheets[0]
    assert [c.value for c in next(ws.iter_rows(max_row=1))] == \
        [s.title for s in SPEC_SETS]
