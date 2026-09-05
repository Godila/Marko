"""Выгрузной артефакт для 1С: GET /v1/nkmt/batches/{id}/report?format=xlsx|csv.

Только published-карточки батча; xlsx — лист «GTIN», csv — BOM для Excel.
"""
import io

import openpyxl
import pytest

from marko.nkmt.models import Batch, Card
from tests.test_api_nkmt_dicts import AUTH_RO, client  # noqa: F401  (фикстура client)


@pytest.fixture
def seeds_pub(db):
    """Батч published с published- и error-карточками (в отчёт — только первая)."""
    b = Batch(source_filename="pub.xlsx", status="published")
    db.add(b); db.flush()
    db.add(Card(article="GH-1", gtin="4630520699970", batch_id=b.id,
                tnved="6109100000", name="Шапка", status="published"))
    db.add(Card(article="GH-2", gtin="", batch_id=b.id,
                tnved="6109100000", name="Брак", status="error"))
    db.commit()
    return b.id


def test_report_xlsx_and_csv(db, client, seeds_pub):   # seeds_pub: батч с published + error карточками
    r = client.get(f"/v1/nkmt/batches/{seeds_pub}/report", headers=AUTH_RO)
    assert r.status_code == 200 and r.content[:2] == b"PK" and r.headers["content-type"].startswith("application/vnd")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("GTIN", "Наименование") and len(rows) == 2  # только published

    r2 = client.get(f"/v1/nkmt/batches/{seeds_pub}/report?format=csv", headers=AUTH_RO)
    assert "4630520699970;Шапка" in r2.text


def test_report_missing_batch_404(db, client):
    assert client.get("/v1/nkmt/batches/99999/report", headers=AUTH_RO).status_code == 404
    assert client.get("/v1/nkmt/batches/99999/report?format=csv", headers=AUTH_RO).status_code == 404


def test_report_bad_format_400(db, client, seeds_pub):
    assert client.get(f"/v1/nkmt/batches/{seeds_pub}/report?format=pdf",
                      headers=AUTH_RO).status_code == 400


def test_report_default_and_headers(db, client, seeds_pub):
    r = client.get(f"/v1/nkmt/batches/{seeds_pub}/report", headers=AUTH_RO)  # без format → xlsx
    assert r.content[:2] == b"PK"
    assert r.headers["content-disposition"] == f'attachment; filename="nkmt-batch-{seeds_pub}.xlsx"'
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.active.title == "GTIN"  # лист называется «GTIN»
    assert list(wb.active.iter_rows(values_only=True)) == [("GTIN", "Наименование"),
                                                            ("4630520699970", "Шапка")]
    r2 = client.get(f"/v1/nkmt/batches/{seeds_pub}/report?format=csv", headers=AUTH_RO)
    assert r2.headers["content-type"].startswith("text/csv")
    assert r2.headers["content-disposition"] == f'attachment; filename="nkmt-batch-{seeds_pub}.csv"'
    assert r2.content.startswith(b"\xef\xbb\xbf")  # BOM — Excel-совместимо
    assert r2.text.splitlines()[0] == "\ufeffGTIN;Наименование"  # BOM + шапка
