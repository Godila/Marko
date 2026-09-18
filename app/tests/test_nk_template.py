"""Шаблон выгрузки: parse.SPEC — единый источник для шапки, констант и инструкции."""
import io

import openpyxl


def test_spec_derived_constants():
    from marko.nkmt.parse import COLUMNS, DEFAULTED_KEYS, REQUIRED_ROW_KEYS, SPEC
    assert REQUIRED_ROW_KEYS == ["article", "name", "tnved", "product_type", "color",
                                 "composition", "size"]
    assert DEFAULTED_KEYS == ["brand", "product_type", "size", "size_system",
                              "target_gender", "declaration_number", "declaration_date",
                              "producer", "country"]
    for s in SPEC:
        if s.title:
            assert COLUMNS[s.title.casefold()] == s.key
    assert set(COLUMNS.values()) == {s.key for s in SPEC if s.title}
    # все колонки шаблона файлочитаемы: producer/country больше не «только из дефолтов»
    assert {"producer", "country"} <= set(COLUMNS.values())


def test_build_template():
    from marko.nkmt.parse import SPEC
    from marko.nkmt.template import build_template
    wb = openpyxl.load_workbook(io.BytesIO(build_template()))
    assert wb.sheetnames == ["Выгрузка", "Инструкция"]
    ws = wb["Выгрузка"]
    cols = [s.key for s in SPEC]           # все SPEC-колонки теперь с title
    titles = [s.title for s in SPEC]
    assert [c.value for c in ws[1]] == titles
    assert ws.cell(row=2, column=cols.index("article") + 1).value == "AB-1001"
    # строки-примеры учат валидным значениям справочников ЧЗ: пол — литералы
    # («ЖЕНСКИЙ», «УНИВЕРСАЛЬНЫЙ (УНИСЕКС)»), размерная система, у шапок
    # размер — из справочника 46–62 (обхват головы)
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    assert rows[0][cols.index("size")] == "M"
    assert rows[0][cols.index("target_gender")] == "ЖЕНСКИЙ"
    assert rows[0][cols.index("size_system")] == "МЕЖДУНАРОДНЫЙ"
    assert rows[1][cols.index("product_type")] == "ШАПКА"
    assert rows[1][cols.index("size")] == "54"
    assert rows[1][cols.index("size_system")] == "ОБХВАТ ГОЛОВЫ"
    assert rows[1][cols.index("target_gender")] == "УНИВЕРСАЛЬНЫЙ (УНИСЕКС)"
    info = wb["Инструкция"]
    assert info.cell(row=1, column=1).value == "Колонка"
    assert info.max_row >= len(titles)   # по строке на колонку + абзацы правил


def test_template_formatting():
    """Читабельность: закреплённая шапка, жирные заголовки, ширины колонок,
    перенос описаний в «Инструкции». Парсер понимает собственный шаблон."""
    from marko.nkmt.parse import SPEC, parse_xlsx
    from marko.nkmt.template import build_template
    data = build_template()
    ws = openpyxl.load_workbook(io.BytesIO(data))["Выгрузка"]
    assert ws.freeze_panes == "A2" and ws["A1"].font.bold
    for i in range(1, len(SPEC) + 1):
        assert (ws.column_dimensions[chr(64 + i)].width or 0) >= 12
    info = openpyxl.load_workbook(io.BytesIO(data))["Инструкция"]
    assert info.freeze_panes == "A2" and info["A1"].font.bold
    assert info["D2"].alignment.wrap_text is True
    notes = [str(r[0]) for r in info.iter_rows(min_row=2, values_only=True)]
    assert "Техрегламент" in notes and "Подстановка" in notes
    parsed = parse_xlsx(data)
    assert len(parsed) == 2 and parsed[1]["size"] == "54"
    assert parsed[1]["target_gender"] == "УНИВЕРСАЛЬНЫЙ (УНИСЕКС)"


def test_template_hint_lists_gender_dictionary():
    """Инструкция должна назвать литералы справочника ЧЗ: инцидент 17.09 —
    оператор писал «Унисекс»/«Универсальный», а в справочнике только
    составной литерал «УНИВЕРСАЛЬНЫЙ (УНИСЕКС)»."""
    from marko.nkmt.parse import SPEC
    hint = next(s.hint for s in SPEC if s.key == "target_gender")
    assert "УНИВЕРСАЛЬНЫЙ (УНИСЕКС)" in hint and "синоним" in hint
